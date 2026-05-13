import docx
import os

docx_path = r'c:\Users\quan2\Downloads\chatbot\Đề Cương Đồ Án tốt nghiệp  (Đã sửa chữa).docx'
out_path = r'c:\Users\quan2\Downloads\chatbot\Đề Cương Đồ Án tốt nghiệp (Đã thêm phần BCE).docx'

doc = docx.Document(docx_path)

# Insert after a specific point or just append to the end.
doc.add_heading('2.5. Phân loại các lớp tham gia ca sử dụng (Mô hình BCE)', level=2)

doc.add_paragraph('Trong quá trình thiết kế hệ thống, các lớp được phân chia theo mô hình Boundary-Control-Entity (BCE) để đảm bảo tính tách biệt, dễ bảo trì và dễ mở rộng. Dưới đây là bảng phân loại các lớp tham gia vào các ca sử dụng chính của hệ thống:')

table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'STT'
hdr_cells[1].text = 'Tên lớp'
hdr_cells[2].text = 'Phân loại (BCE)'
hdr_cells[3].text = 'Vai trò trong hệ thống / Ca sử dụng'

data = [
    ('1', 'View (React Components)', 'Boundary', 'Giao diện trực tiếp tương tác với người dùng (nhập câu hỏi, xem lịch sử). Tham gia mọi ca sử dụng.'),
    ('2', 'API Controller', 'Boundary', 'Nhận HTTP request từ Frontend và trả về response. Giao tiếp giữa người dùng và logic nội bộ.'),
    ('3', 'AuthService / UserService', 'Control', 'Xử lý logic nghiệp vụ đăng nhập, quản lý thông tin tài khoản.'),
    ('4', 'RAGService / ChatbotService', 'Control', 'Điều phối luồng xử lý RAG, gọi LLM, truy xuất thông tin từ Vector DB (Ca sử dụng Hỏi-Đáp).'),
    ('5', 'DatabaseRepository', 'Control', 'Xử lý các logic truy vấn (SELECT, INSERT) vào database (Ca sử dụng Tìm kiếm lịch sử, Thống kê).'),
    ('6', 'User / Role', 'Entity', 'Lưu trữ trạng thái người dùng (Ca sử dụng Quản lý tài khoản, Đăng nhập).'),
    ('7', 'Conversation / Message', 'Entity', 'Lưu trữ thông tin lịch sử chat (Ca sử dụng Tìm kiếm lịch sử chat).'),
    ('8', 'Document / Chunk', 'Entity', 'Lưu trữ tài liệu và đoạn văn bản cho AI RAG.'),
    ('9', 'TokenUsage / Statistics', 'Entity', 'Lưu trữ số liệu dùng token (Ca sử dụng Thống kê & Báo cáo).')
]

for stt, name, bce, role in data:
    row_cells = table.add_row().cells
    row_cells[0].text = stt
    row_cells[1].text = name
    row_cells[2].text = bce
    row_cells[3].text = role

doc.add_paragraph('\nVí dụ chi tiết: Các lớp tham gia Ca sử dụng "Tìm kiếm lịch sử chat"')
p_example = doc.add_paragraph()
r1 = p_example.add_run('- Lớp Boundary (Biên): ')
r1.bold = True
p_example.add_run('Giao diện ô tìm kiếm ở Frontend (CustomSider.jsx) nhận từ khóa từ người dùng. Lớp API Controller nhận request tìm kiếm lịch sử chat.\n')

r2 = p_example.add_run('- Lớp Control (Điều khiển): ')
r2.bold = True
p_example.add_run('Các hàm xử lý logic (như get_conversations) tiến hành kiểm tra ID người dùng, lọc từ khóa, và điều phối lấy dữ liệu.\n')

r3 = p_example.add_run('- Lớp Entity (Thực thể): ')
r3.bold = True
p_example.add_run('Bảng dữ liệu Conversation và Message cung cấp dữ liệu được lưu trữ trong cơ sở dữ liệu để trả về kết quả.')

doc.save(out_path)
print(f"File saved to {out_path}")
